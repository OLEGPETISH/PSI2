import logging
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ticket_generator.models import Assignment, Student

logger = logging.getLogger(__name__)
RESULT_HEADERS = (
    "Группа",
    "Фамилия",
    "Имя",
    "Номер билета",
    "Дата и время",
    "Повтор",
)


class JournalError(Exception):
    """The result journal cannot be read or has an unexpected format."""


class ResultJournal:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if not path.exists():
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Результаты"
            worksheet.append(RESULT_HEADERS)
            try:
                workbook.save(path)
            finally:
                workbook.close()
        workbook = self._open()
        workbook.close()

    def _open(self):
        try:
            workbook = load_workbook(self.path)
        except (InvalidFileException, OSError, zipfile.BadZipFile) as error:
            raise JournalError(f"Не удалось открыть журнал: {error}") from error
        headers = next(
            workbook.active.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        if headers[: len(RESULT_HEADERS)] != RESULT_HEADERS:
            workbook.close()
            raise JournalError(
                "Некорректные заголовки results.xlsx; ожидаются: "
                + ", ".join(RESULT_HEADERS)
            )
        return workbook

    def find_first_ticket(self, student: Student) -> int | None:
        workbook = self._open()
        try:
            for row_number, row in enumerate(
                workbook.active.iter_rows(min_row=2, values_only=True), start=2
            ):
                if row[0:3] != (student.group, student.surname, student.first_name):
                    continue
                value = row[3] if len(row) > 3 else None
                try:
                    number = int(value)
                except (TypeError, ValueError):
                    logger.warning("Некорректный билет в строке %d: %r", row_number, value)
                    continue
                if isinstance(value, float) and not value.is_integer():
                    logger.warning("Дробный билет в строке %d: %r", row_number, value)
                    continue
                return number
            return None
        finally:
            workbook.close()

    def append(self, student: Student, assignment: Assignment) -> None:
        workbook = self._open()
        try:
            workbook.active.append(
                (
                    student.group,
                    student.surname,
                    student.first_name,
                    assignment.ticket_number,
                    datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "да" if assignment.is_repeat else "нет",
                )
            )
            workbook.save(self.path)
        finally:
            workbook.close()
