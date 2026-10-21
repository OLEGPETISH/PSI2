import logging
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from ticket_generator.models import Student, Ticket

logger = logging.getLogger(__name__)
_TICKET_HEADING = re.compile(r"^Билет\s+(\d+)\s*$", re.IGNORECASE)
_QUESTION = re.compile(r"^\d+\.\s*(.+?)\s*$")


class InputDataError(Exception):
    """An input file is missing, unreadable, or contains no usable data."""


def load_students(path: Path) -> dict[str, list[Student]]:
    if not path.is_file():
        raise InputDataError(f"Файл со студентами не найден: {path}")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, OSError, zipfile.BadZipFile) as error:
        raise InputDataError(f"Не удалось открыть файл студентов: {error}") from error

    groups: dict[str, list[Student]] = {}
    try:
        for sheet in workbook.worksheets:
            group = sheet.title.strip()
            students: list[Student] = []
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True),
                start=2,
            ):
                surname = str(row[0]).strip() if row[0] is not None else ""
                first_name = str(row[1]).strip() if row[1] is not None else ""
                if not surname and not first_name:
                    continue
                if not surname or not first_name:
                    logger.warning(
                        "Пропущена неполная запись: группа %r, строка %d",
                        group,
                        row_number,
                    )
                    continue
                students.append(Student(group, surname, first_name))
            groups[group] = students
    finally:
        workbook.close()
    return groups


def load_tickets(path: Path) -> dict[int, Ticket]:
    if not path.is_file():
        raise InputDataError(f"Файл с билетами не найден: {path}")
    try:
        document = Document(path)
    except (PackageNotFoundError, OSError, zipfile.BadZipFile, ValueError) as error:
        raise InputDataError(f"Не удалось открыть файл билетов: {error}") from error

    tickets: dict[int, Ticket] = {}
    current_number: int | None = None
    questions: list[str] = []

    def finish_ticket() -> None:
        if current_number is None:
            return
        if len(questions) < 3:
            logger.warning(
                "Пропущен билет № %d: вопросов %d вместо 3",
                current_number,
                len(questions),
            )
            return
        if current_number in tickets:
            logger.warning("Пропущен повторный билет № %d", current_number)
            return
        tickets[current_number] = Ticket(current_number, tuple(questions[:3]))

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        heading = _TICKET_HEADING.fullmatch(text)
        if heading:
            finish_ticket()
            current_number = int(heading.group(1))
            questions = []
            continue
        if text.casefold().startswith("билет"):
            finish_ticket()
            logger.warning("Нераспознанный заголовок билета: %r", text)
            current_number = None
            questions = []
            continue
        question = _QUESTION.fullmatch(text)
        if current_number is not None and question:
            questions.append(question.group(1).strip())
    finish_ticket()
    if not tickets:
        raise InputDataError("В документе нет билетов минимум с тремя вопросами.")
    return tickets
