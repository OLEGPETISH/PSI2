from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from ticket_generator.journal import JournalError, RESULT_HEADERS, ResultJournal
from ticket_generator.models import Assignment, Student


def test_journal_appends_rows_and_keeps_original_ticket(tmp_path: Path) -> None:
    journal = ResultJournal(tmp_path / "results.xlsx")
    student = Student("ИС-21", "Иванов", "Иван")
    journal.append(student, Assignment(4, False))
    journal.append(student, Assignment(4, True))

    workbook = load_workbook(journal.path, data_only=True)
    rows = list(workbook.active.iter_rows(values_only=True))
    workbook.close()
    assert rows[0] == RESULT_HEADERS
    assert len(rows) == 3
    assert rows[1][:4] == ("ИС-21", "Иванов", "Иван", 4)
    assert rows[1][5] == "нет"
    assert rows[2][:4] == ("ИС-21", "Иванов", "Иван", 4)
    assert rows[2][5] == "да"


def test_first_matching_row_is_authoritative(tmp_path: Path) -> None:
    journal = ResultJournal(tmp_path / "results.xlsx")
    student = Student("ИС-21", "Иванов", "Иван")
    journal.append(student, Assignment(3, False))
    journal.append(student, Assignment(8, True))

    assert journal.find_first_ticket(student) == 3


def test_empty_existing_workbook_reports_journal_error(tmp_path: Path) -> None:
    path = tmp_path / "results.xlsx"
    workbook = Workbook()
    workbook.save(path)
    workbook.close()

    with pytest.raises(JournalError, match="заголовки"):
        ResultJournal(path)
