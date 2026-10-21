from pathlib import Path

from docx import Document
from openpyxl import Workbook
import pytest

from ticket_generator.data_sources import InputDataError, load_students, load_tickets


def test_load_students_maps_each_sheet_to_a_group(tmp_path: Path) -> None:
    path = tmp_path / "students.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ИС-21"
    sheet.append(("Фамилия", "Имя"))
    sheet.append(("Иванов", "Иван"))
    workbook.create_sheet("Пустая группа")
    workbook.save(path)
    workbook.close()

    groups = load_students(path)

    assert groups["ИС-21"][0].display_name == "Иванов Иван"
    assert groups["ИС-21"][0].group == "ИС-21"
    assert groups["Пустая группа"] == []


def test_load_tickets_skips_invalid_heading_and_short_ticket(
    tmp_path: Path,
) -> None:
    path = tmp_path / "tickets.docx"
    document = Document()
    document.add_paragraph("Билет 1")
    document.add_paragraph("1. Только один вопрос")
    document.add_paragraph("Билет неверно")
    document.add_paragraph("Билет 7")
    document.add_paragraph("1. Первый вопрос")
    document.add_paragraph("2. Второй вопрос")
    document.add_paragraph("3. Третий вопрос")
    document.add_paragraph("4. Лишний вопрос")
    document.save(path)

    tickets = load_tickets(path)

    assert list(tickets) == [7]
    assert tickets[7].questions == ("Первый вопрос", "Второй вопрос", "Третий вопрос")


def test_missing_source_is_reported_as_input_error(tmp_path: Path) -> None:
    with pytest.raises(InputDataError, match="не найден"):
        load_students(tmp_path / "missing.xlsx")
